"""文档加载器测试（真实解析 txt / md / docx）。"""
import pytest

from app.utils.loaders import load_document


def test_load_txt(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("你好，世界", encoding="utf-8")
    assert load_document(path) == "你好，世界"


def test_load_markdown(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("# 标题\n\n正文内容", encoding="utf-8")
    assert "正文内容" in load_document(path)


def test_load_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("第一段")
    doc.add_paragraph("第二段")
    path = tmp_path / "a.docx"
    doc.save(path)

    text = load_document(path)

    assert "第一段" in text
    assert "第二段" in text


def test_load_unsupported_type_raises(tmp_path):
    path = tmp_path / "a.bin"
    path.write_bytes(b"\x00\x01")
    with pytest.raises(ValueError):
        load_document(path)


def test_load_pdf(tmp_path):
    import fitz

    path = tmp_path / "a.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Hello PDF Content")
    doc.save(path)
    doc.close()

    assert "Hello PDF Content" in load_document(path)
